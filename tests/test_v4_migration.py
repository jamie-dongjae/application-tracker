"""v3 -> v4 migration: defaults derived from legacy status, new sheets created."""
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook

from tracker.excel.store import ExcelStore


def _v3_workbook(path, rows):
    """Minimal v3-layout workbook: Applications + Meta only."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append(["ID", "Company", "Job Title", "Status", "Date Applied", "Last Updated"])
    for row in rows:
        ws.append(row)
    meta = wb.create_sheet("Meta")
    meta.append(["schema_version", 3])
    meta.append(["next_app_id", len(rows) + 1])
    meta.append(["next_prep_id", 1])
    wb.save(path)


def test_v3_migrates_with_derived_defaults(tmp_path):
    path = tmp_path / "tracker.xlsx"
    old = (date.today() - timedelta(days=45)).isoformat()
    fresh = (date.today() - timedelta(days=3)).isoformat()
    _v3_workbook(path, [
        [1, "Acme", "Analyst", "Rejected", "2026-06-01", ""],
        [2, "Beta", "Engineer", "Interview", "2026-07-01", ""],
        [3, "Gamma", "PM", "Applied", old, f"{old}T09:00:00"],
        [4, "Delta", "Designer", "Applied", fresh, f"{fresh}T09:00:00"],
        [5, "Epsilon", "Lead", "Withdrawn", "2026-06-15", ""],
    ])

    store = ExcelStore(path)
    rows = {r["id"]: r for r in store.list_applications()}

    assert rows[1]["track"] == "career"
    assert rows[1]["current_state"] == "closed"
    assert rows[1]["outcome"] == "rejected_cv"
    assert rows[1]["closed_by"] == "employer"

    assert rows[2]["stage_reached"] == "recruiter_screen"
    assert rows[2]["current_state"] == "awaiting_response"

    assert rows[3]["current_state"] == "stale"       # 45 days silent
    assert rows[4]["current_state"] == "awaiting_response"

    assert rows[5]["outcome"] == "withdrawn_by_me"
    assert rows[5]["closed_by"] == "me"

    # statuses unchanged by the migration
    assert [rows[i]["status"] for i in range(1, 6)] == \
        ["Rejected", "Interview", "Applied", "Applied", "Withdrawn"]

    # persisted: v4 sheets + version bump
    wb = load_workbook(path)
    assert "Events" in wb.sheetnames and "Employers" in wb.sheetnames
    meta = dict(row[:2] for row in wb["Meta"].iter_rows(values_only=True))
    assert meta["schema_version"] == 4


def test_events_and_employers_roundtrip(store):
    rec = store.add_application({"company": "Acme", "title": "Analyst"})
    store.add_event(rec["id"], {"date": "2026-09-01", "event": "recruiter_screen", "note": "call"})
    store.add_event(rec["id"], {"date": "2026-08-20", "event": "applied", "note": ""})
    events = store.list_events(rec["id"])
    assert [e["event"] for e in events] == ["applied", "recruiter_screen"]  # date-sorted

    store.replace_employers([{"employer": "Acme", "rule": "one at a time", "status": "slot free"}])
    fresh = ExcelStore(store.path)
    assert [e["event"] for e in fresh.list_events(rec["id"])] == ["applied", "recruiter_screen"]
    assert fresh.list_employers()[0]["employer"] == "Acme"

    # delete cascades events
    fresh.delete_application(rec["id"])
    assert fresh.list_events(rec["id"]) == []


def test_events_api_and_undo(client):
    rec = client.post("/api/applications", json={"company": "Acme", "title": "Analyst"}).json()
    ev = client.post(f"/api/applications/{rec['id']}/events",
                     json={"date": "2026-09-01", "event": "note", "note": "hello"}).json()
    events = client.get(f"/api/applications/{rec['id']}/events").json()["events"]
    assert len(events) == 1 and events[0]["note"] == "hello"

    client.post("/api/undo")  # undo event creation
    assert client.get(f"/api/applications/{rec['id']}/events").json()["events"] == []

    ev = client.post(f"/api/applications/{rec['id']}/events",
                     json={"event": "nudge"}).json()
    client.delete(f"/api/events/{ev['id']}")
    assert client.get(f"/api/applications/{rec['id']}/events").json()["events"] == []
    client.post("/api/undo")  # undo event deletion restores it
    assert [e["event"] for e in client.get(f"/api/applications/{rec['id']}/events").json()["events"]] == ["nudge"]


def test_employers_api(client):
    body = [{"employer": "KPMG", "rule": "one active", "status": "slot in use"}]
    resp = client.put("/api/employers", json=body).json()
    assert resp["employers"][0]["employer"] == "KPMG"
    assert client.get("/api/employers").json()["employers"][0]["rule"] == "one active"


def test_patch_v4_fields_via_api(client):
    rec = client.post("/api/applications", json={"company": "Acme", "title": "Analyst"}).json()
    patched = client.patch(f"/api/applications/{rec['id']}",
                           json={"stage_reached": "hiring_manager"}).json()
    assert patched["status"] == "Interview"
    assert client.patch(f"/api/applications/{rec['id']}",
                        json={"track": "bogus"}).status_code == 422
